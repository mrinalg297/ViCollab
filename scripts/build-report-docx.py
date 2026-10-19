# Builds a styled .docx from a Markdown file in docs/ (needs pandoc and python-docx).
# Usage: python3 scripts/build-report-docx.py [TEST_REPORT | UI_DESIGN]
import subprocess, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

name = sys.argv[1] if len(sys.argv) > 1 else 'TEST_REPORT'
subprocess.run(['pandoc', f'{name}.md', '-o', f'{name}.docx', '--resource-path=.', '--no-highlight'], cwd='docs', check=True)
d = Document(f'docs/{name}.docx')

# landscape A4 so the wide test case table fits
for s in d.sections:
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.6)

d.styles['Normal'].font.name = 'Calibri'
d.styles['Normal'].font.size = Pt(10.5)
for style, size, color in [('Title', 24, '0E1233'), ('Heading 1', 17, '3040FF'), ('Heading 2', 14, '0E1233'), ('Heading 3', 12, '0E1233')]:
    if style in d.styles:
        f = d.styles[style].font
        f.size, f.color.rgb, f.name = Pt(size), RGBColor.from_string(color), 'Calibri'

def add_borders(tbl):
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'B8BCD0')
        b.append(el)
    tbl._tbl.tblPr.append(b)

def shade(cell, fill):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(sh)

for t in d.tables:
    add_borders(t)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            if i == 0:
                shade(c, 'E3E6F5')
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i == 0:
                        r.bold = True
                    if r.text.strip() in ('Pass', 'Fixed, retest passed'):
                        r.font.color.rgb = RGBColor(0x14, 0x96, 0x6A)

for p in d.paragraphs:
    for r in p.runs:
        if p.style.name == 'Source Code' or (r.style and r.style.name == 'Verbatim Char'):
            r.font.size = Pt(8)

# keep screenshots inside the page
for s in d.inline_shapes:
    ratio = s.height / s.width
    w = min(s.width, Inches(6.2)); h = int(w * ratio)
    if h > Inches(6.3):
        h = Inches(6.3); w = int(h / ratio)
    s.width, s.height = int(w), int(h)

d.save(f'docs/{name}.docx')
print(f'docs/{name}.docx written:', len(d.tables), 'tables,', len(d.inline_shapes), 'images')
